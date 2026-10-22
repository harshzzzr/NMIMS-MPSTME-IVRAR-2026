import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Styles
    NAVY = RGBColor(16, 44, 87)
    TEAL = RGBColor(14, 116, 144)
    DARK_GRAY = RGBColor(51, 65, 85)
    LIGHT_BG = "F1F5F9"
    HEADER_BG = "0F172A"

    # --- Title Page / Header ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("SVKM's NMIMS MPSTME, MUMBAI\nDEPARTMENT OF INFORMATION TECHNOLOGY & COMPUTER ENGINEERING")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = TEAL

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("IVRAR GROUP 01: VR HOME CINEMA & SPATIAL AUDIO SYSTEM\nOPERATIONAL MANUAL & CONTROLLER MAPPING GUIDE")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Course: Introduction to Virtual Reality & Augmented Reality (702COI002) | Semester V | 2026–2027\nStandard: CEDIA/CTA-RP22 Residential Home Cinema Acoustics & Optics")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- Student Team Table ---
    team_table = doc.add_table(rows=1, cols=4)
    team_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    team_table.autofit = False

    headers = ["Roll No", "Student Name", "Degree Programme", "Technical Focus Area"]
    widths = [Inches(1.0), Inches(1.8), Inches(1.8), Inches(2.2)]

    hdr_cells = team_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEADER_BG)
        set_cell_margins(hdr_cells[i], 120, 120, 150, 150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    team_data = [
        ("I066", "Meahaan Sharma", "MBA (Tech) IT", "Lead XR Systems Architect & OpenXR Rig"),
        ("C034", "Amogh Gupta", "B.Tech Computer (Int.)", "Spatial Acoustics, Raycasting & RT60"),
        ("N083", "Harshit Rai", "MBA (Tech) Computer", "Human Factors, Sightlines & Usability"),
        ("N087", "Rannbir Sachdeva", "MBA (Tech) Computer", "Techno-Managerial Systems Evaluation")
    ]

    for row_idx, data in enumerate(team_data):
        row_cells = team_table.add_row().cells
        bg_col = "FFFFFF" if row_idx % 2 == 0 else LIGHT_BG
        for i, val in enumerate(data):
            row_cells[i].text = val
            set_cell_background(row_cells[i], bg_col)
            set_cell_margins(row_cells[i], 80, 80, 120, 120)
            p = row_cells[i].paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # --- Helper to add styled headings ---
    def add_section_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = "Calibri"
            r.font.bold = True
            if level == 1:
                r.font.size = Pt(15)
                r.font.color.rgb = NAVY
            elif level == 2:
                r.font.size = Pt(13)
                r.font.color.rgb = TEAL
            else:
                r.font.size = Pt(11)
                r.font.color.rgb = DARK_GRAY
        return h

    # --- Helper to add an image with caption ---
    def add_figure(img_path, caption_title, caption_text, width=Inches(6.2)):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            run = p_img.add_run()
            run.add_picture(img_path, width=width)

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            r_num = p_cap.add_run(caption_title + ": ")
            r_num.font.bold = True
            r_num.font.size = Pt(9)
            r_num.font.color.rgb = NAVY

            r_txt = p_cap.add_run(caption_text)
            r_txt.font.italic = True
            r_txt.font.size = Pt(9)
            r_txt.font.color.rgb = DARK_GRAY
        else:
            p_err = doc.add_paragraph(f"[Image not found: {img_path}]")
            p_err.runs[0].font.color.rgb = RGBColor(200, 0, 0)

    # ==========================================
    # SECTION 1: SYSTEM OVERVIEW & ARCHITECTURE
    # ==========================================
    add_section_heading("1. Executive Summary & Virtual Environment Overview")
    doc.add_paragraph(
        "The IVRAR Group 01 Virtual Reality Home Cinema is a comprehensive, scientific simulation designed in Unity 6 "
        "and optimized for Meta Quest 2 and the Meta XR Simulator. Built in strict accordance with the CEDIA/CTA-RP22 "
        "Residential Home Cinema Acoustic and Optical Performance Standards, the system features a 5m (W) x 7m (L) x 3m (H) "
        "acoustically treated screening room, a 16:9 cinema proscenium screen, a ceiling projector assembly, VIP cinema seating, "
        "and real-time 3D spatialized audio synchronized with high-definition video playback."
    )
    doc.add_paragraph(
        "Users can experience true spatial audio attenuation with logarithmic distance rolloff originating directly from the center "
        "of the cinema screen, interact with real-time Sabine/Eyring acoustic reverberation (RT60) telemetry, and seamlessly select and switch "
        "between different MP4 video presentations using either VR controller buttons, an in-world 3D console tablet, or desktop hotkeys."
    )

    # Architecture diagram
    arch_img = r"d:\NMIMS-MPSTME-IVRAR-2026\Group_01_Meahaan_Amogh_Harshit_Rannbir\docs\figures\figure1_system_architecture.png"
    add_figure(arch_img, "Figure 1", "End-to-End System Architecture: Monte Carlo Raycasting, OpenXR Rig, Spatial Audio & CEDIA/CTA-RP22 Telemetry.")

    # ==========================================
    # SECTION 2: SCREENSHOT GALLERY & VR VISUALS
    # ==========================================
    add_section_heading("2. Visual Showcase & High-Resolution Project Screenshots")
    doc.add_paragraph(
        "Below are high-resolution captures of the virtual cinema environment taken directly within the Unity Editor and Meta XR Simulator runtime:"
    )

    ss_dir = r"d:\NMIMS-MPSTME-IVRAR-2026\Group_01_Meahaan_Amogh_Harshit_Rannbir\Documentation\Screenshots"

    # Screenshot 1: VIP Seating POV
    add_figure(
        os.path.join(ss_dir, "screenshot_01_vip_seating_pov.png"),
        "Figure 2",
        "Primary Listening Position (PLP) First-Person View from the VIP Cinema Seat, showcasing the 16:9 cinema screen, stage, and side interactive selector."
    )

    # Screenshot 2: Media Selector Kiosk
    add_figure(
        os.path.join(ss_dir, "screenshot_02_media_selector_kiosk.png"),
        "Figure 3",
        "Interactive 3D Media Selector Console and Side Pedestal located within arm's reach of the VIP chair for seamless video selection."
    )

    # Screenshot 3: Cinema Screen Playback
    add_figure(
        os.path.join(ss_dir, "screenshot_03_cinema_screen_playback.png"),
        "Figure 4",
        "Cinema Screen Front Proscenium View displaying unlit, double-sided high-definition video playback with synchronized 3D spatial audio."
    )

    # Screenshot 4: Wide Architectural Overview
    add_figure(
        os.path.join(ss_dir, "screenshot_04_wide_theater_acoustics.png"),
        "Figure 5",
        "Wide Architectural Overview illustrating the acoustic boundary surfaces (drywall, concrete, glass, hardwood floor, acoustic plaster ceiling)."
    )

    # Screenshot 5: Rear Projector
    add_figure(
        os.path.join(ss_dir, "screenshot_05_rear_projector_and_seating.png"),
        "Figure 6",
        "Rear Theater Perspective highlighting the ceiling-mounted optical projector, VIP seating tier, and surround sound array."
    )

    # Telemetry Figure
    telemetry_img = r"d:\NMIMS-MPSTME-IVRAR-2026\Group_01_Meahaan_Amogh_Harshit_Rannbir\docs\figures\figure2_kinematic_telemetry.png"
    add_figure(telemetry_img, "Figure 7", "Live Schroeder Backward Integration RT60 Impulse Response & Sightline Clearance Telemetry Charts.")

    # ==========================================
    # SECTION 3: HOW THE SYSTEM FUNCTIONS
    # ==========================================
    add_section_heading("3. Technical Engineering Architecture: How It Functions")

    add_section_heading("3.1 Universal Video Playback Pipeline", level=2)
    doc.add_paragraph(
        "The cinema presentation is driven by MovieScreenPlayer.cs, an optimized universal playback controller. "
        "The system operates through the following stages:"
    )
    doc.add_paragraph(
        "1. Dynamic Render Pipeline Binding: Video frames are decoded by Unity's VideoPlayer engine into a dedicated 1920x1080 "
        "RenderTexture (RT_CinemaScreen). This texture is assigned to an active Universal Render Pipeline (URP) Unlit material instance.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "2. Double-Sided Mesh Geometry: A 3D Quad primitive with forward-facing normal (-Z) is mounted at (0, 1.6, 3.41). "
        "To prevent backface culling artifacts from any viewing angle, the material explicitly sets _Cull = 0 (double-sided rendering), "
        "guaranteeing vivid image clarity across the entire seating area.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "3. Synchronized 3D Spatial Audio: Video audio tracks (AAC/MP3) are extracted by VideoAudioOutputMode.AudioSource and fed directly "
        "to a physical AudioSource located at the exact geometric center of the cinema screen. Spatial blend is locked to 1.0 (100% 3D spatial), "
        "with a logarithmic rolloff curve (Min Distance = 1.0m, Max Distance = 25.0m). When the viewer turns their head or navigates the room, "
        "the audio dynamically pans and attenuates identically to a real commercial theater.",
        style='List Bullet'
    )

    add_section_heading("3.2 Dynamic StreamingAssets Auto-Discovery", level=2)
    doc.add_paragraph(
        "To enable seamless operation without requiring Unity Editor recompilation or manual scene wiring, the system scans "
        "Assets/StreamingAssets/ at startup and runtime for any valid media containers (.mp4, .webm, .mov). "
        "All discovered files are indexed into an active playlist with zero-latency switching capability."
    )

    add_section_heading("3.3 Meta XR Building Blocks Architecture", level=2)
    doc.add_paragraph(
        "The project integrates official Meta XR Building Blocks for maximum fidelity on Meta Quest 2 and Meta XR Simulator:"
    )
    doc.add_paragraph(
        "• Camera Rig Building Block (ID: e47682b9-c270-40b1-b16d-90b627a5ce1b): Configures the OVRCameraRig with FloorLevel tracking origin.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Controller Tracking Building Blocks (ID: 5817f7c0-f2a5-45f9-a5ca-64264e0166e8): Enables 6-DOF tracking of Left and Right Touch controllers.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Controller Buttons Mapper (ID: 80aa2eb2-af2f-40a9-ad9e-18044ce65bee): Binds hardware inputs on both controllers directly to cinema playback and playlist navigation routines.",
        style='List Bullet'
    )

    # ==========================================
    # SECTION 4: CONTROLLER BUTTON FUNCTIONS
    # ==========================================
    add_section_heading("4. Controller Button Mappings & Function Reference")
    doc.add_paragraph(
        "Hardware button bindings have been engineered specifically for ergonomic, two-handed operation in VR. "
        "The Right Controller manages video playback and volume, while the Left Controller handles playlist navigation and menus."
    )

    # Table for Right Controller
    add_section_heading("4.1 Right Controller (Playback & Volume Controls)", level=2)
    rc_table = doc.add_table(rows=1, cols=4)
    rc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rc_headers = ["Physical Button", "OVRInput Binding", "Assigned Action", "Operational Description"]
    for i, h in enumerate(rc_headers):
        rc_table.rows[0].cells[i].text = h
        set_cell_background(rc_table.rows[0].cells[i], "1E293B")
        set_cell_margins(rc_table.rows[0].cells[i], 100, 100, 120, 120)
        p = rc_table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    rc_data = [
        ("Button A", "Button.One", "Next Video Clip", "Advances to the next video in the StreamingAssets playlist. Loops back to beginning if at the end."),
        ("Button B", "Button.Two", "Play / Pause Video", "Instantly pauses or resumes playback of both video stream and 3D spatial audio track."),
        ("Thumbstick UP", "SecondaryThumbstickUp", "Volume Up (+10%)", "Increases theater master volume by 10% (clamps at 100%)."),
        ("Thumbstick DOWN", "SecondaryThumbstickDown", "Volume Down (-10%)", "Decreases theater master volume by 10% (clamps at 0%)."),
        ("Index Trigger", "SecondaryIndexTrigger", "Restart Video", "Rewinds video timestamp to 00:00:00 and immediately restarts playback.")
    ]

    for idx, row in enumerate(rc_data):
        cells = rc_table.add_row().cells
        bg = "FFFFFF" if idx % 2 == 0 else LIGHT_BG
        for i, val in enumerate(row):
            cells[i].text = val
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], 80, 80, 100, 100)
            p = cells[i].paragraphs[0]
            if i in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Table for Left Controller
    add_section_heading("4.2 Left Controller (Playlist Navigation & Interface)", level=2)
    lc_table = doc.add_table(rows=1, cols=4)
    lc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lc_headers = ["Physical Button", "OVRInput Binding", "Assigned Action", "Operational Description"]
    for i, h in enumerate(lc_headers):
        lc_table.rows[0].cells[i].text = h
        set_cell_background(lc_table.rows[0].cells[i], "1E293B")
        set_cell_margins(lc_table.rows[0].cells[i], 100, 100, 120, 120)
        p = lc_table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    lc_data = [
        ("Button X", "Button.Three", "Previous Video Clip", "Jumps backward to the previous video clip in the playlist."),
        ("Button Y", "Button.Four", "Toggle Media Selector", "Toggles the in-world / HUD Media Selector Console overlay on and off."),
        ("Menu Button", "Button.Start", "Rescan Video Folder", "Re-scans Assets/StreamingAssets/ to immediately detect newly copied video files."),
        ("Thumbstick LEFT", "PrimaryThumbstickLeft", "Previous Video", "Alternative ergonomic thumb navigation for previous video."),
        ("Thumbstick RIGHT", "PrimaryThumbstickRight", "Next Video", "Alternative ergonomic thumb navigation for next video.")
    ]

    for idx, row in enumerate(lc_data):
        cells = lc_table.add_row().cells
        bg = "FFFFFF" if idx % 2 == 0 else LIGHT_BG
        for i, val in enumerate(row):
            cells[i].text = val
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], 80, 80, 100, 100)
            p = cells[i].paragraphs[0]
            if i in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Table for Simulator / Keyboard Controls
    add_section_heading("4.3 Meta XR Simulator & Desktop Keyboard Hotkeys", level=2)
    doc.add_paragraph(
        "For development, grading, and testing without wearing a headset, the system provides full desktop and simulator hotkeys:"
    )

    kb_table = doc.add_table(rows=1, cols=3)
    kb_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    kb_headers = ["Key / Input", "Subsystem Action", "Description"]
    for i, h in enumerate(kb_headers):
        kb_table.rows[0].cells[i].text = h
        set_cell_background(kb_table.rows[0].cells[i], "0F766E")
        set_cell_margins(kb_table.rows[0].cells[i], 100, 100, 120, 120)
        p = kb_table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    kb_data = [
        ("V", "Toggle Media Selector HUD", "Opens/closes the clickable desktop Media Selector window in Game View."),
        ("Spacebar", "Play / Pause", "Instantly pauses or resumes the current cinema movie."),
        ("] (Right Bracket) / PageDown", "Next Video", "Switches to the next video clip in Assets/StreamingAssets."),
        ("[ (Left Bracket) / PageUp", "Previous Video", "Switches to the previous video clip in Assets/StreamingAssets."),
        ("Tab / H", "Toggle Telemetry HUD", "Displays CEDIA RP22 compliance, RT60 metrics, SPL readings, and optics telemetry."),
        ("1 / 2 / 3", "Switch Acoustic Scenarios", "1 = Baseline Bare Room (RT60: 0.72s), 2 = Over-Treated (RT60: 0.18s), 3 = CEDIA Optimized (RT60: 0.38s)."),
        ("M", "Run Acoustic Raytracing", "Fires 1,000 Monte Carlo acoustic rays across the cinema boundary meshes."),
        ("R", "Reset Cinema & Telemetry", "Restarts movie to 00:00 and resets room acoustic materials to baseline."),
        ("WASD + Right Mouse", "First-Person Navigation", "Moves the virtual listening position through the home cinema.")
    ]

    for idx, row in enumerate(kb_data):
        cells = kb_table.add_row().cells
        bg = "FFFFFF" if idx % 2 == 0 else LIGHT_BG
        for i, val in enumerate(row):
            cells[i].text = val
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], 80, 80, 100, 100)
            p = cells[i].paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ==========================================
    # SECTION 5: STEP-BY-STEP OPERATOR GUIDE
    # ==========================================
    add_section_heading("5. Operator Guide: Adding & Switching Videos")

    add_section_heading("5.1 Where and How to Save Videos", level=2)
    doc.add_paragraph(
        "To add new movies or presentations to the cinema system:",
        style='List Bullet'
    )
    doc.add_paragraph(
        "1. Open Windows File Explorer and navigate to:\n"
        "   D:\\NMIMS-MPSTME-IVRAR-2026\\Group_01_Meahaan_Amogh_Harshit_Rannbir\\Assets\\StreamingAssets\\",
        style='List Bullet'
    )
    doc.add_paragraph(
        "2. Copy and paste any standard .mp4, .webm, or .mov file directly into this folder. "
        "Any descriptive filename is valid (e.g., movie.mp4, dolby_trailer.mp4, imax_demo.mp4).",
        style='List Bullet'
    )
    doc.add_paragraph(
        "3. Supported Encoding Standards:\n"
        "   • Video Codec: H.264 / AVC (Fast Start enabled, non-fragmented container).\n"
        "   • Audio Codec: AAC-LC or MP3 (Stereo 2.0 or 5.1 Surround Sound).\n"
        "   • Recommended Resolution: 1080p (1920x1080) at 24, 30, or 60 FPS.",
        style='List Bullet'
    )

    add_section_heading("5.2 How to Switch Videos During Simulation", level=2)
    doc.add_paragraph(
        "Method A (VR Controllers): Press Button A on the Right Controller to cycle forward, or Button X on the Left Controller to cycle backward.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Method B (Interactive Media Selector HUD): Press V on your keyboard. In the bottom-left window, click any video title directly with your mouse.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Method C (In-World 3D Console): In the virtual cinema room, look directly to the left of the VIP cinema seat at (-1.15, 0.85, -0.30). "
        "The 3D tablet displays the active playlist and allows direct pointer selection.",
        style='List Bullet'
    )

    # Save Document
    out_path = r"d:\NMIMS-MPSTME-IVRAR-2026\Group_01_Meahaan_Amogh_Harshit_Rannbir\Documentation\IVRAR_Group_01_Home_Cinema_VR_Systems_Manual.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_document()
